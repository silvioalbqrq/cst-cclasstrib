# -*- coding: utf-8 -*-
"""Gera data/anexo-viii.json a partir da planilha do Anexo VIII
(Correlação Item NBS / IndOp / cClassTrib — IBS/CBS).

Uso:
    python scripts/atualizar_anexo8.py "caminho/AnexoVIII....xlsx"

A planilha usa células mescladas verticalmente: cada NBS se estende
pelas linhas-variante de IndOp (030101-030104 = local da prestação;
100301 = Domicílio principal do adquirente, com PS Onerosa / Adq.
Exterior). Este script "desmescla" a grade e gera um registro por
linha, isto é, uma correlação NBS x IndOp x cClassTrib por registro.

Normalizações:
- IndOp armazenado como número perde zeros à esquerda -> zfill(6)
- PS Onerosa / Adq. Exterior: "s"/"n" -> "S"/"N"
- Local de incidência: capitalização unificada
"""
import json, re, sys
import openpyxl

DEFAULT_XLSX = "AnexoVIII-CorrelacaoItemNBSIndOpCClassTrib_IBSCBS_V1.00.00.xlsx"
OUT = "data/anexo-viii.json"

LOCAIS = {
    "local da entrega ou disponibilização": "Local da entrega ou disponibilização",
    "local da prestação": "Local da prestação",
    "local do evento": "Local do evento",
    "local evento": "Local do evento",
    "local do imóvel": "Local do imóvel",
    "via explorada": "Via explorada",
    "domicílio principal do adquirente": "Domicílio principal do adquirente",
}


def clean(v):
    return "" if v is None else str(v).strip()


def norm_indop(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return str(int(v)).zfill(6)
    return clean(v).zfill(6) if clean(v).isdigit() else clean(v)


def norm_local(v):
    key = clean(v).lower()
    return LOCAIS.get(key, clean(v))


def unmerge(ws):
    """Expande células mescladas: devolve {(linha, coluna): valor}."""
    grid = {}
    for row in ws.iter_rows():
        for cell in row:
            grid[(cell.row, cell.column)] = cell.value
    for m in ws.merged_cells.ranges:
        v = grid.get((m.min_row, m.min_col))
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                grid[(r, c)] = v
    return grid


def main(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # ---------- matriz REGRA inc. X ----------
    # leitura direta (sem desmesclar): E:F é merge horizontal visual,
    # o valor de N/S (coluna F) é vazio de propósito
    rg = wb["REGRA inc. X"]

    def g(r, c):
        return clean(rg.cell(row=r, column=c).value)

    # linhas 3-5 = dados; colunas B/C = PS onerosa SIM, E/F = PS onerosa NÃO
    regra = [
        {"psOnerosa": "S", "adqExterior": "N",
         "indop": [x for x in (g(3, 2), g(4, 2), g(5, 2)) if x]},
        {"psOnerosa": "S", "adqExterior": "S",
         "indop": [x for x in (g(3, 3), g(4, 3), g(5, 3)) if x]},
        {"psOnerosa": "N", "adqExterior": "N",
         "indop": [x for x in (g(3, 5), g(4, 5), g(5, 5)) if x]},
        {"psOnerosa": "N", "adqExterior": "S",
         "indop": [x for x in (g(3, 6), g(4, 6), g(5, 6)) if x]},
    ]

    # ---------- aba tabela geral ----------
    ws = wb["tabela geral"]
    grid = unmerge(ws)

    itens = []
    for r in range(2, ws.max_row + 1):
        vals = [clean(grid.get((r, c))) for c in range(1, 11)]
        item, desc_item, nbs, desc_nbs, ps, adq, indop, local, cct, nome_cct = vals
        indop = norm_indop(indop)
        if not (indop or nbs or cct):
            continue  # linha espaçadora
        nbs_digits = re.sub(r"\D", "", nbs)
        itens.append({
            "itemLc116": item,
            "descItem": desc_item,
            "nbs": nbs,
            "nbsDigits": nbs_digits,
            "descNbs": desc_nbs,
            "psOnerosa": ps.upper(),
            "adqExterior": adq.upper(),
            "indop": indop,
            "localIncidencia": norm_local(local),
            "cClassTrib": cct,
            "nomeClassTrib": nome_cct,
        })

    # remove duplicatas exatas (células mescladas podem gerar repetições)
    vistas = set()
    unicos = []
    for it in itens:
        chave = json.dumps(it, ensure_ascii=False)
        if chave not in vistas:
            vistas.add(chave)
            unicos.append(it)
    itens = unicos

    # ---------- validação ----------
    problemas = []
    for i, it in enumerate(itens):
        if it["nbs"] and len(it["nbsDigits"]) != 9:
            problemas.append(f"registro {i}: NBS '{it['nbs']}' não tem 9 dígitos")
        if it["indop"] and len(it["indop"]) != 6:
            problemas.append(f"registro {i}: IndOp '{it['indop']}' não tem 6 dígitos")
    if problemas:
        print("PROBLEMAS:")
        for p in problemas[:20]:
            print("  " + p)
        raise SystemExit(1)

    nbs_unicos = sorted({it["nbsDigits"] for it in itens if it["nbsDigits"]})
    cct_unicos = sorted({it["cClassTrib"] for it in itens if it["cClassTrib"]})
    indops = sorted({it["indop"] for it in itens if it["indop"]})
    locais = sorted({it["localIncidencia"] for it in itens if it["localIncidencia"]})

    data = {
        "meta": {
            "fonte": "Anexo VIII — Correlação Item NBS / IndOp / cClassTrib — IBS/CBS",
            "observacao": (
                "Extraído da planilha com desmesclagem de células: cada registro é uma "
                "correlação NBS x IndOp x cClassTrib. IndOp numérico foi normalizado "
                "com zeros à esquerda (6 dígitos). PS Onerosa/Adq. Exterior vazios "
                "aplicam-se às regras de incidência por local de prestação (0301xx)."
            ),
            "geradoEm": "2026-10-08",
            "correlacoes": len(itens),
            "nbsUnicos": len(nbs_unicos),
            "cClassTribUnicos": len(cct_unicos),
            "indops": indops,
            "locaisIncidencia": locais,
        },
        "regraIndOp": regra,
        "itens": itens,
    }

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    print(f"OK -> {OUT}")
    print(f"  correlações: {len(itens)}")
    print(f"  NBS unicos: {len(nbs_unicos)} | cClassTrib unicos: {len(cct_unicos)} | IndOps: {len(indops)}")
    print(f"  matriz regra: {json.dumps(regra, ensure_ascii=False)}")
    print(f"  locais: {locais}")


if __name__ == "__main__":
    caminho = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    main(caminho)
